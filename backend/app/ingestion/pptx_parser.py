from pptx import Presentation
from typing import List, Dict, Any

class PPTXParser:
    @staticmethod
    def parse(file_path: str) -> List[Dict[str, Any]]:
        prs = Presentation(file_path)
        slides = []
        
        for idx, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)
                elif shape.has_table:
                    table_rows = []
                    for row in shape.table.rows:
                        table_rows.append(" | ".join([cell.text.strip() for cell in row.cells]))
                    slide_texts.append("\n".join(table_rows))
            
            slides.append({
                "page_number": idx + 1,
                "text": "\n".join(slide_texts),
                "sections": [{"title": f"Slide {idx+1}", "offset": 0}],
                "bbox": [0, 0, 1024, 768]
            })
            
        return slides
